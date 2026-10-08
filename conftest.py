from __future__ import annotations

import pytest


def pytest_addoption(parser: pytest.Parser) -> None:
    parser.addoption("--browser", choices=["chrome", "edge"], default="chrome", help="Visible browser")
    parser.addoption("--pause", type=float, default=5.0, help="Seconds to keep each result visible")


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item: pytest.Item, call: pytest.CallInfo):
    """Keep the Excel status column in sync with the latest executed case."""
    outcome = yield
    report = outcome.get_result()
    if report.when != "call":
        return
    case = getattr(item, "callspec", None)
    case = case.params.get("case") if case else None
    if not case:
        return

    from test_login import CASE_FILE
    from openpyxl import load_workbook

    workbook = load_workbook(CASE_FILE)
    sheet = workbook["Login"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == case["Test Case ID"]:
            sheet.cell(row, 10).value = "Passed" if report.passed else "Failed"
            break
    workbook.save(CASE_FILE)
