"""Excel-driven, visible-browser Selenium tests for the UTC login page."""

from __future__ import annotations

import argparse
import os
import sys
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


ROOT = Path(__file__).resolve().parent
CASE_FILE = ROOT / "test_cases" / "login_test_cases.xlsx"
BASE_URL = os.getenv("LOGIN_URL", "https://vanphongdientu.utc.edu.vn/Login")
HEADERS = [
    "Test Case ID",
    "Module",
    "Test Case",
    "Preconditions",
    "Username",
    "Password",
    "Steps",
    "Expected Result",
    "Action",
    "Status",
]


def ensure_workbook() -> None:
    CASE_FILE.parent.mkdir(parents=True, exist_ok=True)
    if not CASE_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Login"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:J1"
        workbook.save(CASE_FILE)


def get_cases():
    ensure_workbook()
    sheet = load_workbook(CASE_FILE, data_only=True).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    return [dict(zip(HEADERS, row)) for row in rows if row[0]]


def set_case_status(case_id: str, status: str) -> None:
    workbook = load_workbook(CASE_FILE)
    sheet = workbook["Login"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == case_id:
            sheet.cell(row, 10).value = status
            break
    workbook.save(CASE_FILE)


def make_driver(browser: str):
    if browser == "edge":
        options = webdriver.EdgeOptions()
        options.add_argument("--start-maximized")
        return webdriver.Edge(options=options)
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    return webdriver.Chrome(options=options)


def run_case(case: dict, browser: str, pause_seconds: float) -> None:
    driver = make_driver(browser)
    wait = WebDriverWait(driver, 20)
    case_id = case["Test Case ID"]
    try:
        driver.get(BASE_URL)
        username = wait.until(EC.visibility_of_element_located((By.NAME, "username")))
        password = wait.until(EC.visibility_of_element_located((By.NAME, "userpwd")))
        submit = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "input.submit_login")))
        assert "Đăng nhập" in driver.title, f"Unexpected page title: {driver.title!r}"

        action = case["Action"]
        if action != "page_load":
            if case["Username"] is not None:
                username.send_keys(str(case["Username"]))
            if case["Password"] is not None:
                password.send_keys(str(case["Password"]))
            submit.click()
            try:
                wait.until(lambda current: current.execute_script("return document.readyState") == "complete")
            except TimeoutException:
                pass

        # These cases check the login UI and that a rejected/empty login does not
        # navigate away into the authenticated application.
        if action == "page_load":
            assert username.is_displayed() and password.is_displayed() and submit.is_displayed()
            assert "Đăng nhập" in submit.get_attribute("value")
        else:
            wait.until(EC.presence_of_element_located((By.NAME, "username")))
            assert "/Login" in driver.current_url, f"Unexpected navigation: {driver.current_url}"

        print(f"PASS {case_id}: {case['Test Case']}")
        if pause_seconds > 0:
            print(f"Giữ cửa sổ {pause_seconds:g} giây để quan sát...")
            time.sleep(pause_seconds)
    except Exception:
        artifact_dir = ROOT / "artifacts"
        artifact_dir.mkdir(exist_ok=True)
        screenshot = artifact_dir / f"{case_id}.png"
        driver.save_screenshot(str(screenshot))
        print(f"FAIL {case_id}; screenshot: {screenshot}", file=sys.stderr)
        raise
    finally:
        driver.quit()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Run visible Selenium login cases from Excel")
    parser.add_argument("--case", help="Run one test case ID; default runs all cases")
    parser.add_argument("--browser", choices=["chrome", "edge"], default="chrome")
    parser.add_argument("--pause", type=float, default=2.0, help="Seconds to leave each result visible")
    args = parser.parse_args()
    cases = get_cases()
    selected = [case for case in cases if not args.case or case["Test Case ID"] == args.case]
    if not selected:
        print(f"Không tìm thấy test case: {args.case or '(file Excel chưa có case)'}", file=sys.stderr)
        return 2
    failed = 0
    for case in selected:
        try:
            run_case(case, args.browser, args.pause)
            set_case_status(case["Test Case ID"], "Passed")
        except Exception as error:
            failed += 1
            set_case_status(case["Test Case ID"], "Failed")
            print(f"{type(error).__name__}: {error}", file=sys.stderr)
    print(f"Kết quả: {len(selected) - failed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
