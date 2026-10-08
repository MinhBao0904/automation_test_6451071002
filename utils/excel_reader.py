"""Load and validate test cases from the project Excel workbook."""

from __future__ import annotations

from openpyxl import load_workbook

from config.settings import TEST_CASE_FILE


HEADERS = (
    "Test Case ID", "Module", "Test Case", "Preconditions", "Username",
    "Password", "Steps", "Expected Result", "Action", "Status",
)


def load_test_cases() -> list[dict]:
    if not TEST_CASE_FILE.is_file():
        raise FileNotFoundError(f"Không tìm thấy file test case: {TEST_CASE_FILE}")

    workbook = load_workbook(TEST_CASE_FILE, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(min_row=2, values_only=True)
        cases = [dict(zip(HEADERS, row)) for row in rows if row and row[0]]
    finally:
        workbook.close()
    if not cases:
        raise ValueError(f"Workbook chưa có test case: {TEST_CASE_FILE}")

    case_ids = [case["Test Case ID"] for case in cases]
    if len(case_ids) != len(set(case_ids)):
        raise ValueError("Test Case ID phải duy nhất trong file Excel")

    for case in cases:
        for column in ("Test Case", "Steps", "Expected Result", "Action"):
            if not case.get(column):
                raise ValueError(f"{case['Test Case ID']}: thiếu dữ liệu cột {column}")
    return cases


def set_case_statuses(statuses: dict[str, str]) -> None:
    workbook = load_workbook(TEST_CASE_FILE)
    sheet = workbook["Login"]
    for row in range(2, sheet.max_row + 1):
        case_id = sheet.cell(row, 1).value
        if case_id in statuses:
            sheet.cell(row, 10).value = statuses[case_id]
    workbook.save(TEST_CASE_FILE)
